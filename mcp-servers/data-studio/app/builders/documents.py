"""Relatórios a partir de markdown: DOCX (Word), HTML estilizado ou .md. Blocos especiais no markdown:
`[[table: SELECT ...]]` vira uma tabela com o resultado da consulta; `[[image: figura_1.png]]` insere a imagem."""
import base64
import re
from pathlib import Path

import markdown as md

from .common import filename, fmt, table_data

BLOCK = re.compile(r"^\[\[(table|image):\s*(.+?)\]\]\s*$")

HTML_CSS = """body{font-family:Inter,system-ui,Segoe UI,sans-serif;max-width:860px;margin:40px auto;padding:0 20px;
color:#1f2937;line-height:1.6}h1{color:#4f46e5}h2{border-bottom:2px solid #e5e7eb;padding-bottom:4px;margin-top:32px}
table{border-collapse:collapse;width:100%;margin:12px 0;font-size:14px}th{background:#4f46e5;color:#fff;text-align:left}
th,td{padding:6px 10px;border-bottom:1px solid #e5e7eb}img{max-width:100%}code{background:#f3f4f6;padding:1px 5px;border-radius:4px}
blockquote{border-left:4px solid #4f46e5;margin:0;padding:4px 16px;color:#4b5563;background:#f9fafb}"""


def _expand_md(text: str, base: Path, q) -> str:
    out = []
    for line in text.splitlines():
        m = BLOCK.match(line.strip())
        if m and m.group(1) == "table":
            cols, rows = table_data({"sql": m.group(2)}, q, limit=200)
            out.append("| " + " | ".join(cols) + " |")
            out.append("|" + "---|" * len(cols))
            out += ["| " + " | ".join(fmt(v) if isinstance(v, (int, float)) else str(v) for v in r) + " |" for r in rows]
        elif m and m.group(1) == "image":
            p = base / m.group(2).strip()
            if p.exists():
                b64 = base64.b64encode(p.read_bytes()).decode()
                out.append(f"![{p.stem}](data:image/png;base64,{b64})")
        else:
            out.append(line)
    return "\n".join(out)


def _docx(title: str, text: str, base: Path, q, path: Path):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    h = d.add_heading(title, 0)
    h.runs[0].font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)

    def runs(par, s):
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", s):
            if part.startswith("**") and part.endswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = "Consolas"
            elif part:
                par.add_run(part)

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        m = BLOCK.match(line.strip())
        if m:
            if m.group(1) == "table":
                cols, rows = table_data({"sql": m.group(2)}, q, limit=200)
                _table(d, cols, [[fmt(v) if isinstance(v, (int, float)) else str(v) for v in r] for r in rows])
            else:
                p = base / m.group(2).strip()
                if p.exists():
                    d.add_picture(str(p), width=Inches(6))
        elif line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1]):
            head = [c.strip() for c in line.strip("|").split("|")]
            rows, i = [], i + 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            _table(d, head, rows)
            continue
        elif line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 4)
            d.add_heading(line.lstrip("#").strip(), level)
        elif re.match(r"^\s*[-*] ", line):
            runs(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+[.)] ", line):
            runs(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)] ", "", line))
        elif line.startswith(">"):
            p = d.add_paragraph(style="Intense Quote")
            runs(p, line.lstrip("> "))
        elif line.strip():
            runs(d.add_paragraph(), line)
        i += 1
    d.save(path)


def _table(d, cols, rows):
    t = d.add_table(rows=1, cols=len(cols))
    t.style = "Light Grid Accent 1"
    for j, c in enumerate(cols):
        t.rows[0].cells[j].text = str(c)
    for r in rows:
        cells = t.add_row().cells
        for j, v in enumerate(r[:len(cols)]):
            cells[j].text = str(v)


def build(spec: dict, base: Path, q) -> str:
    title = spec.get("title", "Relatório")
    fmt_ = spec.get("format", "docx")
    text = spec.get("markdown", "")
    stem = spec.get("filename") or title
    if fmt_ == "docx":
        name = filename(stem, ".docx")
        _docx(title, text, base, q, base / name)
    elif fmt_ == "md":
        name = filename(stem, ".md")
        (base / name).write_text(f"# {title}\n\n" + _expand_md(text, base, q), encoding="utf-8")
    else:
        name = filename(stem, ".html")
        body = md.markdown(_expand_md(text, base, q), extensions=["tables", "fenced_code"])
        (base / name).write_text(f"<!doctype html><html lang=pt-BR><head><meta charset=utf-8><title>{title}</title>"
                                 f"<style>{HTML_CSS}</style></head><body><h1>{title}</h1>{body}</body></html>",
                                 encoding="utf-8")
    return name
