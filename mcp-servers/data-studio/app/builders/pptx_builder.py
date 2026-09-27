"""PowerPoint 16:9 com gráficos NATIVOS (editáveis no PowerPoint), tabelas, KPIs, imagens e notas do apresentador."""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .common import chart_data, filename, fmt, scalar, table_data, theme

W, H = Inches(13.333), Inches(7.5)
KINDS = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
         "stacked": XL_CHART_TYPE.COLUMN_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA,
         "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "scatter": XL_CHART_TYPE.XY_SCATTER}


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self, spec, query_df):
        self.spec, self.q, self.t = spec, query_df, theme(spec)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # ---------------------------------------------------------------- base
    def slide(self, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(self.t["accent"] if dark else self.t["bg"])
        return s

    def text(self, s, x, y, w, h, text, size=18, bold=False, color=None, align=PP_ALIGN.LEFT):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = str(text)
        r.font.size, r.font.bold = Pt(size), bold
        r.font.color.rgb = rgb(color or self.t["fg"])
        return tb

    def header(self, s, title):
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.55), Inches(0.12), Inches(0.62))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(self.t["accent"])
        bar.line.fill.background()
        self.text(s, Inches(0.9), Inches(0.42), Inches(11.8), Inches(0.9), title, 30, True)

    def footer(self, s, sl):
        foot = self.spec.get("footer") or self.spec.get("title", "")
        self.text(s, Inches(0.6), Inches(7.0), Inches(9), Inches(0.35), foot, 10, color=self.t["muted"])
        self.text(s, Inches(11.9), Inches(7.0), Inches(0.9), Inches(0.35), str(self.n), 10,
                  color=self.t["muted"], align=PP_ALIGN.RIGHT)
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]

    # ---------------------------------------------------------------- tipos
    def title(self, sl):
        s = self.slide(dark=True)
        self.text(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6), sl.get("title") or self.spec.get("title", ""),
                  44, True, "FFFFFF")
        if sl.get("subtitle") or self.spec.get("subtitle"):
            self.text(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.0),
                      sl.get("subtitle") or self.spec.get("subtitle"), 22, color="E0E7FF")
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]

    def section(self, sl):
        s = self.slide(dark=True)
        self.text(s, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.4), sl.get("title", ""), 38, True, "FFFFFF")
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]

    def bullets(self, sl, x=Inches(0.9), w=Inches(11.6), items=None, s=None):
        own = s is None
        if own:
            s = self.slide()
            self.header(s, sl.get("title", ""))
        tb = s.shapes.add_textbox(x, Inches(1.6), w, Inches(5.2))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items if items is not None else sl.get("bullets", [])):
            level = 1 if str(item).startswith(("  ", "- ")) else 0
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.level = level
            r = p.add_run()
            r.text = ("• " if level == 0 else "– ") + str(item).lstrip(" -")
            r.font.size = Pt(22 if level == 0 else 18)
            r.font.color.rgb = rgb(self.t["fg"] if level == 0 else self.t["muted"])
            p.space_after = Pt(10)
        if own:
            self.footer(s, sl)

    def two_columns(self, sl):
        s = self.slide()
        self.header(s, sl.get("title", ""))
        for i, side in enumerate(("left", "right")):
            x = Inches(0.9 + i * 6.1)
            if sl.get(f"{side}_title"):
                self.text(s, x, Inches(1.45), Inches(5.6), Inches(0.5), sl[f"{side}_title"], 20, True,
                          self.t["accent"])
            self.bullets(sl, x=x, w=Inches(5.6), items=sl.get(side, []), s=s)
        self.footer(s, sl)

    def text_slide(self, sl):
        s = self.slide()
        self.header(s, sl.get("title", ""))
        self.text(s, Inches(0.9), Inches(1.7), Inches(11.6), Inches(5), sl.get("text", ""), 22)
        self.footer(s, sl)

    def chart(self, sl):
        s = self.slide()
        self.header(s, sl.get("title", ""))
        ch = sl.get("chart") or {}
        kind = ch.get("kind", "column")
        cats, series = chart_data(ch, self.q)
        top = Inches(1.5)
        h = Inches(4.9 if sl.get("caption") else 5.3)
        if kind == "scatter":
            data = XyChartData()
            for se in series:
                sr = data.add_series(se["name"])
                for xv, yv in zip(cats, se["values"], strict=False):
                    try:
                        sr.add_data_point(float(xv), yv or 0)
                    except ValueError:
                        continue
        else:
            data = CategoryChartData()
            data.categories = cats
            for se in (series[:1] if kind in ("pie", "doughnut") else series):
                data.add_series(se["name"], [v if v is not None else 0 for v in se["values"]])
        gf = s.shapes.add_chart(KINDS.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED), Inches(0.9), top, Inches(11.6), h, data)
        c = gf.chart
        # com uma série só, o PowerPoint usa o nome da série como título automático — o título do slide já basta
        c.has_title = False
        c.font.size = Pt(12)
        c.font.color.rgb = rgb(self.t["fg"])
        c.has_legend = len(series) > 1 or kind in ("pie", "doughnut")
        if c.has_legend:
            c.legend.position = XL_LEGEND_POSITION.BOTTOM
            c.legend.include_in_layout = False
        if kind in ("pie", "doughnut"):
            plot = c.plots[0]
            plot.has_data_labels = True
            plot.data_labels.number_format = "0.0%"
            plot.data_labels.show_percentage, plot.data_labels.show_value = True, False
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(self.t["palette"][i % len(self.t["palette"])])
        else:
            for i, sr in enumerate(c.plots[0].series):
                col = rgb(self.t["palette"][i % len(self.t["palette"])])
                if kind in ("line", "scatter"):
                    sr.format.line.color.rgb = col
                    sr.format.line.width = Pt(2.5)
                else:
                    sr.format.fill.solid()
                    sr.format.fill.fore_color.rgb = col
            if ch.get("data_labels") and kind != "scatter":
                c.plots[0].has_data_labels = True
                c.plots[0].data_labels.font.size = Pt(10)
        if sl.get("caption"):
            self.text(s, Inches(0.9), Inches(6.45), Inches(11.6), Inches(0.5), sl["caption"], 14, color=self.t["muted"])
        self.footer(s, sl)

    def table(self, sl):
        s = self.slide()
        self.header(s, sl.get("title", ""))
        cols, rows = table_data(sl.get("table") or sl, self.q, limit=int(sl.get("max_rows", 12)))
        cols, rows = cols[:8], [r[:8] for r in rows]
        if not cols:
            self.footer(s, sl)
            return
        shape = s.shapes.add_table(len(rows) + 1, len(cols), Inches(0.9), Inches(1.55), Inches(11.6),
                                   Emu(int(Inches(0.42)) * (len(rows) + 1)))
        tbl = shape.table
        for j, c in enumerate(cols):
            cell = tbl.cell(0, j)
            cell.text = str(c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(self.t["accent"])
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.bold, para.runs[0].font.size = True, Pt(13)
            para.runs[0].font.color.rgb = rgb("FFFFFF")
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r):
                cell = tbl.cell(i, j)
                cell.text = fmt(v) if isinstance(v, (int, float)) else str(v)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(self.t["card"] if i % 2 else self.t["bg"])
                run = cell.text_frame.paragraphs[0].runs[0]
                run.font.size = Pt(12)
                run.font.color.rgb = rgb(self.t["fg"])
        self.footer(s, sl)

    def kpis(self, sl):
        s = self.slide()
        self.header(s, sl.get("title", ""))
        items = sl.get("items", [])[:6]
        n = max(len(items), 1)
        gap, w = Inches(0.3), int((Inches(11.6) - Inches(0.3) * (n - 1)) / n)
        for i, it in enumerate(items):
            x = Inches(0.9) + i * (w + gap)
            card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), w, Inches(2.8))
            card.fill.solid()
            card.fill.fore_color.rgb = rgb(self.t["card"])
            card.line.fill.background()
            val = fmt(scalar(it, self.q), it.get("format"))
            # fonte proporcional à largura do card: valores longos em cards estreitos não quebram a linha
            size = max(18, min(36, int(w / 914400 * 72 / max(len(val), 1) * 1.7)))
            self.text(s, x, Inches(2.5), w, Inches(1.2), val, size, True, self.t["accent"], PP_ALIGN.CENTER)
            self.text(s, x, Inches(3.7), w, Inches(0.6), it.get("label", ""), 16, color=self.t["fg"],
                      align=PP_ALIGN.CENTER)
            if it.get("delta"):
                d = str(it["delta"])
                self.text(s, x, Inches(4.25), w, Inches(0.5), d, 14, True,
                          "DC2626" if d.strip().startswith("-") else "059669", PP_ALIGN.CENTER)
        self.footer(s, sl)

    def image(self, sl, base: Path):
        s = self.slide()
        self.header(s, sl.get("title", ""))
        p = base / sl.get("path", "")
        if p.exists():
            from PIL import Image
            with Image.open(p) as im:
                ratio = im.width / im.height
            maxw, maxh = Inches(11.6), Inches(4.9)
            w = min(maxw, int(maxh * ratio))
            h = int(w / ratio)
            s.shapes.add_picture(str(p), Inches(0.9) + int((maxw - w) / 2), Inches(1.5), w, h)
        else:
            self.text(s, Inches(0.9), Inches(3), Inches(11.6), Inches(1), f"(imagem não encontrada: {p.name})", 18)
        if sl.get("caption"):
            self.text(s, Inches(0.9), Inches(6.45), Inches(11.6), Inches(0.5), sl["caption"], 14, color=self.t["muted"])
        self.footer(s, sl)


def build(spec: dict, base: Path, stem: str, query_df) -> str:
    d = Deck(spec, query_df)
    slides = spec.get("slides") or []
    if not slides or slides[0].get("type") != "title":
        d.title({})
    for sl in slides:
        kind = sl.get("type", "bullets")
        if kind == "image":
            d.image(sl, base)
        elif kind in ("text", "quote"):
            d.text_slide(sl)
        else:
            getattr(d, kind if kind in ("title", "section", "bullets", "two_columns", "chart", "table", "kpis")
                    else "bullets")(sl)
    name = filename(stem, ".pptx")
    d.prs.save(base / name)
    return name
