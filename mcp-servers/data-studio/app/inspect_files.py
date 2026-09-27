"""Leitura/perfil rápido de qualquer arquivo do workspace: tabelas, PDF, Word, PowerPoint, imagens (OCR), texto."""
import json
from pathlib import Path

import pandas as pd

MAX_TEXT = 15000


def _md(df: pd.DataFrame, n: int) -> str:
    return df.head(n).to_markdown(index=False) if len(df.columns) else "(sem colunas)"


def read_table(p: Path, sheet=None, nrows=None) -> dict[str, pd.DataFrame]:
    s = p.suffix.lower()
    if s in (".csv", ".tsv", ".txt"):
        return {p.stem: pd.read_csv(p, sep=None, engine="python", nrows=nrows, encoding_errors="replace")}
    if s in (".xlsx", ".xlsm", ".xls"):
        sheets = pd.read_excel(p, sheet_name=sheet if sheet is not None else None, nrows=nrows)
        return sheets if isinstance(sheets, dict) else {str(sheet): sheets}
    if s == ".parquet":
        df = pd.read_parquet(p)
        return {p.stem: df.head(nrows) if nrows else df}
    if s in (".json", ".jsonl", ".ndjson"):
        lines = s != ".json"
        return {p.stem: pd.read_json(p, lines=lines, nrows=nrows if lines else None)}
    raise ValueError(f"não é um formato tabular: {s}")


def profile(df: pd.DataFrame, rows: int) -> str:
    info = [f"{len(df):,} linhas × {len(df.columns)} colunas"]
    cols = []
    for c in df.columns:
        s = df[c]
        cols.append(f"| {c} | {s.dtype} | {s.isna().sum()} | {s.nunique(dropna=True)} | "
                    f"{str(s.dropna().iloc[0])[:40] if s.notna().any() else ''} |")
    info.append("| coluna | tipo | nulos | distintos | exemplo |\n|---|---|---|---|---|\n" + "\n".join(cols))
    num = df.select_dtypes("number")
    if len(num.columns):
        info.append("Estatísticas numéricas:\n" + num.describe().T.round(3).to_markdown())
    info.append(f"Primeiras {rows} linhas:\n" + _md(df, rows))
    return "\n\n".join(info)


def inspect(p: Path, max_rows: int = 10, pages: int = 5) -> str:
    s = p.suffix.lower()
    size = p.stat().st_size
    head = f"# {p.name} ({size/1024:.1f} KB)\n"
    if s in (".csv", ".tsv", ".xlsx", ".xlsm", ".xls", ".parquet", ".json", ".jsonl", ".ndjson"):
        tables = read_table(p)
        parts = [head]
        for name, df in tables.items():
            parts.append(f"## Tabela/aba: {name}\n{profile(df, max_rows)}")
        return "\n\n".join(parts)[:MAX_TEXT * 2]
    if s == ".pdf":
        import pdfplumber
        out = [head]
        with pdfplumber.open(p) as pdf:
            out.append(f"{len(pdf.pages)} página(s).")
            for i, page in enumerate(pdf.pages[:pages]):
                text = page.extract_text() or ""
                out.append(f"## Página {i+1}\n{text.strip()}")
                for j, tb in enumerate(page.extract_tables()[:3]):
                    if tb and len(tb) > 1:
                        df = pd.DataFrame(tb[1:], columns=[str(c) for c in tb[0]])
                        out.append(f"### Tabela {j+1} (página {i+1})\n{_md(df, 30)}")
            if len(pdf.pages) > pages:
                out.append(f"… mais {len(pdf.pages) - pages} página(s): use inspect_file com pages maior ou run_python.")
        text = "\n\n".join(out)
        if len(text.strip()) < 200:
            text += "\n\n(Pouco texto extraível — PDF provavelmente escaneado. Use ocr=true em inspect_file.)"
        return text[:MAX_TEXT * 2]
    if s == ".docx":
        import docx
        d = docx.Document(p)
        paras = "\n".join(x.text for x in d.paragraphs if x.text.strip())
        tabs = []
        for i, t in enumerate(d.tables[:5]):
            rows = [[c.text for c in r.cells] for r in t.rows]
            if len(rows) > 1:
                tabs.append(f"### Tabela {i+1}\n" + _md(pd.DataFrame(rows[1:], columns=rows[0]), 30))
        return (head + paras + "\n\n" + "\n\n".join(tabs))[:MAX_TEXT]
    if s == ".pptx":
        from pptx import Presentation
        prs = Presentation(p)
        out = [head, f"{len(prs.slides)} slide(s)."]
        for i, sl in enumerate(prs.slides, 1):
            texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            out.append(f"## Slide {i}\n" + "\n".join(texts))
        return "\n\n".join(out)[:MAX_TEXT]
    if s in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"):
        from PIL import Image
        with Image.open(p) as im:
            return head + f"Imagem {im.width}×{im.height}, modo {im.mode}. Use ocr=true para extrair texto."
    try:
        text = p.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return head + f"(binário não suportado: {e})"
    if s == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except Exception:
            pass
    return head + text[:MAX_TEXT]


def ocr(p: Path, lang: str = "por+eng", pages: int = 5) -> str:
    import pytesseract
    from PIL import Image
    s = p.suffix.lower()
    if s == ".pdf":
        import pdfplumber
        out = []
        with pdfplumber.open(p) as pdf:
            for i, page in enumerate(pdf.pages[:pages]):
                im = page.to_image(resolution=200).original
                out.append(f"## Página {i+1}\n" + pytesseract.image_to_string(im, lang=lang))
        return "\n\n".join(out)[:MAX_TEXT]
    with Image.open(p) as im:
        return pytesseract.image_to_string(im.convert("RGB"), lang=lang)[:MAX_TEXT]
